from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.db.models import Count, Q
from django.utils import timezone
from datetime import timedelta
import json

from core.models import Sector, EstadoDocumento
from entrada.models import Expediente, MovimentacaoDocumento
from users.models import User


@login_required
def dashboard_simples(request):
    """
    Dashboard simples com KPIs principais e gráficos diretos.
    """
    from django.db.models import Avg, F, Q
    from datetime import datetime, timedelta
    
    # Filtros do formulário
    data_inicio = request.GET.get('data_inicio')
    data_fim = request.GET.get('data_fim')
    sector_id = request.GET.get('sector')
    estado_id = request.GET.get('estado')
    prioridade = request.GET.get('prioridade')
    
    # Definir período padrão (últimos 30 dias)
    if not data_inicio:
        data_inicio = (timezone.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    if not data_fim:
        data_fim = timezone.now().strftime('%Y-%m-%d')
    
    # Converter para datetime
    data_inicio_dt = datetime.strptime(data_inicio, '%Y-%m-%d').date()
    data_fim_dt = datetime.strptime(data_fim, '%Y-%m-%d').date()
    
    # Query base com filtros
    query_filtros = Q(ativo=True, data_criacao__date__range=[data_inicio_dt, data_fim_dt])
    
    if sector_id:
        query_filtros &= Q(sector_responsavel_id=sector_id)
    if estado_id:
        query_filtros &= Q(estado_atual_id=estado_id)
    if prioridade:
        query_filtros &= Q(prioridade=prioridade)
    
    # KPIs principais - Performance
    total_documentos = Expediente.objects.filter(query_filtros).count()
    
    # Documentos por estado atual
    documentos_recebidos = Expediente.objects.filter(query_filtros, estado_atual__nome='Recebido').count()
    documentos_encaminhados = Expediente.objects.filter(query_filtros, estado_atual__nome='Encaminhado').count()
    documentos_em_tratamento = Expediente.objects.filter(query_filtros, estado_atual__nome='Em Tratamento').count()
    documentos_concluidos = Expediente.objects.filter(query_filtros, estado_atual__nome='Concluído').count()
    documentos_devolvidos = Expediente.objects.filter(query_filtros, estado_atual__nome='Devolvido').count()
    documentos_arquivados = Expediente.objects.filter(query_filtros, estado_atual__nome='Arquivado').count()
    
    # KPIs de Performance - Tempo
    documentos_com_data_aprovacao = Expediente.objects.filter(
        query_filtros, 
        data_aprovacao__isnull=False
    )
    
    tempo_medio_tratamento = documentos_com_data_aprovacao.aggregate(
        tempo_medio=Avg(F('data_aprovacao') - F('data_criacao'))
    )['tempo_medio']
    
    # Converter para dias se existir
    if tempo_medio_tratamento:
        tempo_medio_dias = tempo_medio_tratamento.days
    else:
        tempo_medio_dias = 0
    
    # Documentos em atraso (mais de 7 dias sem movimento)
    data_limite_atraso = timezone.now() - timedelta(days=7)
    documentos_atraso = Expediente.objects.filter(
        query_filtros,
        data_atualizacao__lt=data_limite_atraso,
        estado_atual__nome__in=['Recebido', 'Encaminhado', 'Em Tratamento']
    ).count()
    
    # Taxa de conclusão
    documentos_periodo = Expediente.objects.filter(query_filtros).count()
    documentos_concluidos_periodo = Expediente.objects.filter(
        query_filtros,
        estado_atual__nome='Concluído'
    ).count()
    
    taxa_conclusao = (documentos_concluidos_periodo / documentos_periodo * 100) if documentos_periodo > 0 else 0
    
    # Movimentações no período
    movimentacoes_periodo = MovimentacaoDocumento.objects.filter(
        data_movimentacao__date__range=[data_inicio_dt, data_fim_dt]
    ).count()
    
    # KPIs de Fluxo
    documentos_por_prioridade = Expediente.objects.filter(query_filtros).values(
        'prioridade'
    ).annotate(
        total=Count('id')
    ).order_by('-total')
    
    documentos_por_origem = Expediente.objects.filter(query_filtros).values(
        'origem'
    ).annotate(
        total=Count('id')
    ).order_by('-total')
    
    # KPIs de Sectores
    documentos_por_sector = Expediente.objects.filter(query_filtros).values(
        'sector_responsavel__nome'
    ).annotate(
        total=Count('id')
    ).order_by('-total')[:10]
    
    # Carga de trabalho por utilizador
    carga_por_utilizador = Expediente.objects.filter(
        query_filtros,
        utilizador_atual__isnull=False
    ).values(
        'utilizador_atual__first_name',
        'utilizador_atual__last_name',
        'utilizador_atual__tipo_utilizador'
    ).annotate(
        total=Count('id')
    ).order_by('-total')[:5]
    
    # Dados para gráficos
    estados_labels = ['Recebido', 'Encaminhado', 'Em Tratamento', 'Concluído', 'Devolvido', 'Arquivado']
    estados_data = [
        documentos_recebidos, documentos_encaminhados, documentos_em_tratamento,
        documentos_concluidos, documentos_devolvidos, documentos_arquivados
    ]
    
    sectores_labels = [item['sector_responsavel__nome'] or 'Sem Sector' for item in documentos_por_sector]
    sectores_data = [item['total'] for item in documentos_por_sector]
    
    prioridades_labels = [item['prioridade'].title() for item in documentos_por_prioridade]
    prioridades_data = [item['total'] for item in documentos_por_prioridade]
    
    origens_labels = [item['origem'].title() for item in documentos_por_origem]
    origens_data = [item['total'] for item in documentos_por_origem]
    
    # Dados da linha temporal (período filtrado)
    documentos_timeline = Expediente.objects.filter(query_filtros).extra(
        select={'dia': 'DATE(data_criacao)'}
    ).values('dia').annotate(
        total=Count('id')
    ).order_by('dia')
    
    timeline_labels = [item['dia'].strftime('%d/%m') for item in documentos_timeline]
    timeline_data = [item['total'] for item in documentos_timeline]
    
    # Opções para filtros
    sectores = Sector.objects.filter(ativo=True).order_by('nome')
    estados = EstadoDocumento.objects.filter(ativo=True).order_by('ordem')
    prioridades_opcoes = Expediente.PRIORIDADES
    
    context = {
        'titulo': 'Dashboard de Relatórios - FTC Expedientes',
        'filtros': {
            'data_inicio': data_inicio,
            'data_fim': data_fim,
            'sector_id': sector_id,
            'estado_id': estado_id,
            'prioridade': prioridade,
            'sectores': sectores,
            'estados': estados,
            'prioridades': prioridades_opcoes,
        },
        'kpis': {
            # Performance
            'total_documentos': total_documentos,
            'tempo_medio_tratamento': tempo_medio_dias,
            'documentos_atraso': documentos_atraso,
            'taxa_conclusao': round(taxa_conclusao, 1),
            'movimentacoes_periodo': movimentacoes_periodo,
            
            # Estados
            'documentos_recebidos': documentos_recebidos,
            'documentos_encaminhados': documentos_encaminhados,
            'documentos_em_tratamento': documentos_em_tratamento,
            'documentos_concluidos': documentos_concluidos,
            'documentos_devolvidos': documentos_devolvidos,
            'documentos_arquivados': documentos_arquivados,
        },
        'graficos': {
            'estados': {
                'labels': json.dumps(estados_labels),
                'data': json.dumps(estados_data)
            },
            'sectores': {
                'labels': json.dumps(sectores_labels),
                'data': json.dumps(sectores_data)
            },
            'prioridades': {
                'labels': json.dumps(prioridades_labels),
                'data': json.dumps(prioridades_data)
            },
            'origens': {
                'labels': json.dumps(origens_labels),
                'data': json.dumps(origens_data)
            },
            'timeline': {
                'labels': json.dumps(timeline_labels),
                'data': json.dumps(timeline_data)
            }
        },
        'tabelas': {
            'carga_utilizadores': carga_por_utilizador,
            'documentos_por_sector': documentos_por_sector,
        }
    }
    
    return render(request, 'relatorios/dashboard_simples.html', context)


@login_required
def exportar_relatorio(request, formato):
    """
    Exporta relatório simples em PDF, Excel ou CSV.
    """
    if formato not in ['pdf', 'excel', 'csv']:
        return JsonResponse({'error': 'Formato inválido'}, status=400)
    
    # Dados básicos
    total_documentos = Expediente.objects.filter(ativo=True).count()
    documentos_pendentes = Expediente.objects.filter(
        ativo=True,
        estado_atual__nome__in=['Recebido', 'Encaminhado', 'Em Tratamento']
    ).count()
    documentos_concluidos = Expediente.objects.filter(
        ativo=True,
        estado_atual__nome='Concluído'
    ).count()
    
    documentos_por_sector = Expediente.objects.filter(ativo=True).values(
        'sector_responsavel__nome'
    ).annotate(
        total=Count('id')
    ).order_by('-total')
    
    documentos_por_estado = Expediente.objects.filter(ativo=True).values(
        'estado_atual__nome'
    ).annotate(
        total=Count('id')
    ).order_by('-total')
    
    if formato == 'pdf':
        return exportar_pdf(total_documentos, documentos_pendentes, documentos_concluidos, 
                           documentos_por_sector, documentos_por_estado)
    elif formato == 'excel':
        return exportar_excel(total_documentos, documentos_pendentes, documentos_concluidos, 
                             documentos_por_sector, documentos_por_estado)
    elif formato == 'csv':
        return exportar_csv(total_documentos, documentos_pendentes, documentos_concluidos, 
                           documentos_por_sector, documentos_por_estado)


def exportar_pdf(total_doc, pendentes, concluidos, por_sector, por_estado):
    """Exporta relatório em PDF."""
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from io import BytesIO
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    
    # Título
    title = Paragraph("Relatório de Documentos - FTC", styles['Title'])
    story.append(title)
    story.append(Spacer(1, 20))
    
    # KPIs
    kpi_data = [
        ['Métrica', 'Valor'],
        ['Total de Documentos', str(total_doc)],
        ['Documentos Pendentes', str(pendentes)],
        ['Documentos Concluídos', str(concluidos)],
    ]
    
    kpi_table = Table(kpi_data)
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    story.append(kpi_table)
    story.append(Spacer(1, 20))
    
    # Documentos por Sector
    sector_data = [['Sector', 'Total de Documentos']]
    for item in por_sector:
        sector_data.append([item['sector_responsavel__nome'] or 'Sem Sector', str(item['total'])])
    
    sector_table = Table(sector_data)
    sector_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    story.append(sector_table)
    
    doc.build(story)
    buffer.seek(0)
    
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="relatorio_documentos.pdf"'
    return response


def exportar_excel(total_doc, pendentes, concluidos, por_sector, por_estado):
    """Exporta relatório em Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from io import BytesIO
    
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório de Documentos"
    
    # KPIs
    ws['A1'] = 'Métrica'
    ws['B1'] = 'Valor'
    ws['A2'] = 'Total de Documentos'
    ws['B2'] = total_doc
    ws['A3'] = 'Documentos Pendentes'
    ws['B3'] = pendentes
    ws['A4'] = 'Documentos Concluídos'
    ws['B4'] = concluidos
    
    # Estilo do cabeçalho
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    for cell in ws['A1:B1']:
        cell.font = header_font
        cell.fill = header_fill
    
    # Documentos por Sector
    ws['A6'] = 'Sector'
    ws['B6'] = 'Total de Documentos'
    
    row = 7
    for item in por_sector:
        ws[f'A{row}'] = item['sector_responsavel__nome'] or 'Sem Sector'
        ws[f'B{row}'] = item['total']
        row += 1
    
    # Estilo do cabeçalho da tabela
    for cell in ws['A6:B6']:
        cell.font = header_font
        cell.fill = header_fill
    
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="relatorio_documentos.xlsx"'
    return response


def exportar_csv(total_doc, pendentes, concluidos, por_sector, por_estado):
    """Exporta relatório em CSV."""
    import csv
    from io import StringIO
    
    buffer = StringIO()
    writer = csv.writer(buffer)
    
    # KPIs
    writer.writerow(['Métrica', 'Valor'])
    writer.writerow(['Total de Documentos', total_doc])
    writer.writerow(['Documentos Pendentes', pendentes])
    writer.writerow(['Documentos Concluídos', concluidos])
    writer.writerow([])  # Linha vazia
    
    # Documentos por Sector
    writer.writerow(['Sector', 'Total de Documentos'])
    for item in por_sector:
        writer.writerow([item['sector_responsavel__nome'] or 'Sem Sector', item['total']])
    
    response = HttpResponse(buffer.getvalue(), content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="relatorio_documentos.csv"'
    return response
